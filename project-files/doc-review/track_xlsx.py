#!/usr/bin/env python3
"""Excel for review (24/10/2026, Michael: "work with Review when sending, in Excel, Word and every document").
Excel has no reliable tracked changes, so the review copy shows them instead:

    python3 track_xlsx.py OLD.xlsx NEW.xlsx OUT.xlsx [--author "מיכאל פרייליך"]

OUT is a copy of NEW where every changed cell is yellow with a comment "היה: <old value>", every new cell
(that was empty) is green, a removed value is red with "נמחק: <old>", and a first sheet "שינויים" lists each
change (sheet, cell, old, new). Send OUT for review and keep NEW as the clean version; the controlled change log
(form 02.03) gets one line per document version.
"""
import sys
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill, Font, Alignment

YEL, GRN, RED = (PatternFill('solid', fgColor=c) for c in ('FFF2A8', 'C6EFCE', 'FFC7CE'))


def norm(v):
    return '' if v is None else str(v).strip()


def track(old_path, new_path, out_path, author):
    old, new = load_workbook(old_path), load_workbook(new_path)
    changes = []
    for ws in new.worksheets:
        if ws.title not in old.sheetnames:
            changes.append((ws.title, '', '', 'גיליון חדש'))
            continue
        wo = old[ws.title]
        rows = max(ws.max_row, wo.max_row); cols = max(ws.max_column, wo.max_column)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                a, b = norm(wo.cell(r, c).value), norm(ws.cell(r, c).value)
                if a == b:
                    continue
                cell = ws.cell(r, c)
                if not a:
                    cell.fill = GRN
                elif not b:
                    cell.fill = RED; cell.comment = Comment('נמחק: ' + a[:500], author)
                else:
                    cell.fill = YEL; cell.comment = Comment('היה: ' + a[:500], author)
                changes.append((ws.title, cell.coordinate, a, b))
    for name in old.sheetnames:
        if name not in new.sheetnames:
            changes.append((name, '', '', 'גיליון נמחק'))
    log = new.create_sheet('שינויים', 0)
    log.sheet_view.rightToLeft = True
    log.append(['גיליון', 'תא', 'היה', 'עכשיו'])
    for c in log[1]:
        c.font = Font(bold=True)
    for row in changes:
        log.append([x if x != '' else '(ריק)' for x in row])
    for col, w in zip('ABCD', (18, 8, 50, 50)):
        log.column_dimensions[col].width = w
    for row in log.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='right')
    new.save(out_path)
    return len(changes)


if __name__ == '__main__':
    a = sys.argv[1:]
    author = a[a.index('--author') + 1] if '--author' in a else 'מיכאל פרייליך'
    print('changes:', track(a[0], a[1], a[2], author))
