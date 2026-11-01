#!/usr/bin/env python3
"""tools/form-fill.py (01/11/2026, Michael: fill the factory's controlled Word form, never build one; Word and PDF;
the logo and header kept). Builds a tiny form with python-docx in a temp dir: a header with the logo image, "form
08.01" and "page 1 of 1", an RTL and an LTR label/value table, a participants table with "ID no." in its header, and a
tracked insertion (w:ins). Runs the tool and checks the values, the rows, the accepted change, the summary, the PDF,
and that the logo / form-number checks really fail when the output loses them.
Usage: python3 tests/harness/form-fill-test.py"""
import copy, importlib.util, io, json, os, re, shutil, struct, subprocess, sys, tempfile, zipfile, zlib
from contextlib import redirect_stdout, redirect_stderr

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

HERE = os.path.dirname(os.path.abspath(__file__))
TOOL = os.path.join(HERE, '..', '..', 'tools', 'form-fill.py')
FORM = '\u05d8\u05d5\u05e4\u05e1 08.01'
PAGE = '\u05e2\u05de\u05d5\u05d3 1 \u05de\u05ea\u05d5\u05da 1'
DATE_L = '\u05ea\u05d0\u05e8\u05d9\u05da \u05d4\u05d4\u05d3\u05e8\u05db\u05d4:'
DUR_L = '\u05de\u05e9\u05da \u05d4\u05d4\u05d3\u05e8\u05db\u05d4:'
LAST_L = '\u05d7\u05ea\u05d9\u05de\u05d4'
ID_H = "\u05de\u05e1' \u05d6\u05d4\u05d5\u05ea"
OLD = '\u05e0\u05d5\u05e9\u05d0 \u05d9\u05e9\u05df'
NEW = '\u05e0\u05d5\u05e9\u05d0 \u05d7\u05d3\u05e9'
fails = []
passed = 0


def check(name, ok, info=''):
    global passed
    if ok:
        passed += 1
        print('  ok   ' + name)
    else:
        fails.append(name)
        print('  FAIL ' + name + ('  ' + str(info)[:400] if info else ''))


def png(path):
    raw = b''.join(b'\x00' + b'\xcc\x22\x22' * 8 for _ in range(8))
    chunk = lambda t, d: struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    with open(path, 'wb') as f:
        f.write(b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', 8, 8, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b''))


def build(path, logo):
    d = docx.Document()
    sec = d.sections[0]
    ht = sec.header.add_table(1, 3, Cm(16))
    ht.cell(0, 0).paragraphs[0].add_run().add_picture(logo, width=Cm(2))
    ht.cell(0, 1).text = FORM
    ht.cell(0, 2).text = PAGE
    # RTL label/value table (bidiVisual): the value is the next cell in document order, shown to the left
    t1 = d.add_table(rows=2, cols=4)
    bidi = OxmlElement('w:bidiVisual')
    t1._tbl.tblPr.append(bidi)
    t1.cell(0, 0).text = DATE_L
    t1.cell(0, 2).text = DUR_L
    t1.cell(1, 2).text = 'x'
    t1.cell(1, 3).text = LAST_L
    # LTR label/value table
    t2 = d.add_table(rows=1, cols=2)
    t2.cell(0, 0).text = 'Trainer:'
    # participants
    t3 = d.add_table(rows=2, cols=4)
    for i, h in enumerate(["\u05de\u05e1''\u05d3", '\u05e9\u05dd', ID_H, LAST_L]):
        t3.cell(0, i).text = h
    # a tracked insertion: the old text sits inside w:ins, which python-docx paragraphs do not see
    p = d.add_paragraph()
    ins = OxmlElement('w:ins')
    ins.set(qn('w:id'), '1')
    ins.set(qn('w:author'), 'x')
    ins.set(qn('w:date'), '2012-09-04T00:00:00Z')
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = OLD
    r.append(t)
    ins.append(r)
    p._p.append(ins)
    d.save(path)


def strip_images(src, dst):
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if re.match(r'word/(header\d*|document)\.xml$', it.filename):
                data = re.sub(rb'<w:drawing>.*?</w:drawing>', b'', data, flags=re.S)
            zout.writestr(it, data)


def run(*args):
    r = subprocess.run([sys.executable, TOOL] + list(args), capture_output=True, text=True, timeout=600)
    last = (r.stdout.strip().splitlines() or ['{}'])[-1]
    try:
        js = json.loads(last)
    except ValueError:
        js = {}
    return r.returncode, js, r.stderr


def load_tool():
    spec = importlib.util.spec_from_file_location('form_fill', TOOL)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def main():
    tmp = tempfile.mkdtemp(prefix='ff-test-')
    try:
        logo = os.path.join(tmp, 'logo.png')
        png(logo)
        tpl = os.path.join(tmp, 'form.docx')
        build(tpl, logo)
        with zipfile.ZipFile(tpl) as z:
            check('the template really holds a w:ins (the test is real)', b'<w:ins ' in z.read('word/document.xml'))
        spec = {
            'cells': [{'label': DATE_L, 'value': '02/11/2026'}, {'label': DUR_L, 'value': '60'}, {'label': 'Trainer', 'value': 'Michael'},
                      {'label': LAST_L, 'value': 'S', 'nth': 0}],
            'table_cells': [{'table': 0, 'row': 1, 'col': 0, 'value': 'exact'}],
            'rows': {'header_contains': ID_H, 'count': 22, 'number_col': 0, 'data': [['A', '1'], ['B', '2']]},
            'replace_text': [{'find': OLD, 'value': NEW}],
            'fit_one_page': True,
        }
        sp = os.path.join(tmp, 'spec.json')
        with open(sp, 'w', encoding='utf-8') as f:
            json.dump(spec, f, ensure_ascii=False)
        out = os.path.join(tmp, 'out', 'filled')
        rc, js, err = run('--template', tpl, '--spec', sp, '--out', out)
        check('runs green, one JSON summary line with ok, docx, pdf, pages, images, form_no', rc == 0 and js.get('ok') is True and all(k in js for k in ('docx', 'pdf', 'pages', 'images', 'form_no')), (rc, js, err))
        check('summary: the logo counted (images 1) and the form number from the header (08.01)', js.get('images') == 1 and js.get('form_no') == '08.01', js)
        d = docx.Document(out + '.docx')
        t1, t2, t3 = d.tables[0], d.tables[1], d.tables[2]
        tcs = lambda row: [''.join(t.text or '' for t in x.iter(qn('w:t'))).strip() for x in row._tr.findall(qn('w:tc'))]
        check('RTL table: the value lands in the cell after the label in document order (shown to its left)', tcs(t1.rows[0])[1] == '02/11/2026' and tcs(t1.rows[0])[3] == '60', tcs(t1.rows[0]))
        check('LTR table: "Trainer:" matched without its colon, value to its right', tcs(t2.rows[0]) == ['Trainer:', 'Michael'], tcs(t2.rows[0]))
        check('a label that is the last cell of its row: the value goes to the cell before it', tcs(t1.rows[1])[2] == 'S' and tcs(t1.rows[1])[3] == LAST_L, tcs(t1.rows[1]))
        check('table_cells: exact placement', tcs(t1.rows[1])[0] == 'exact', tcs(t1.rows[1]))
        rows = t3.rows
        check('rows: the participants table extended to 22 data rows', len(rows) == 23, len(rows))
        check('rows: column 0 numbered 1..22', [tcs(r)[0] for r in rows[1:]] == [str(i) for i in range(1, 23)], [tcs(r)[0] for r in rows[1:4]])
        check('rows: data filled into the other columns, the header kept', tcs(rows[1])[1:3] == ['A', '1'] and tcs(rows[2])[1:3] == ['B', '2'] and tcs(rows[3])[1] == '' and ID_H in tcs(rows[0]), [tcs(r) for r in rows[:4]])
        body = '\n'.join(p.text for p in d.paragraphs)
        with zipfile.ZipFile(out + '.docx') as z:
            xml = z.read('word/document.xml').decode('utf-8')
        check('tracked change accepted: the w:ins text is seen and replaced, no w:ins left', NEW in body and OLD not in xml and '<w:ins ' not in xml, body[-80:])
        if shutil.which('soffice'):
            check('PDF made, one page as the template', js.get('pdf') and os.path.exists(js['pdf']) and js.get('pages') == 1 and js.get('template_pages') == 1, js)
            big = dict(spec, rows={'header_contains': ID_H, 'count': 60, 'number_col': 0})
            with open(sp, 'w', encoding='utf-8') as f:
                json.dump(big, f, ensure_ascii=False)
            rc2, js2, err2 = run('--template', tpl, '--spec', sp, '--out', os.path.join(tmp, 'out', 'big'))
            check('fit_one_page: 60 rows are shrunk step by step and fit one page (or fail with a clear message at 220 twips)',
                  (rc2 == 0 and js2.get('pages') == 1 and js2.get('row_height', 9999) < 9999) or (rc2 == 1 and 'does not fit' in err2), (rc2, js2, err2[-300:]))
            huge = dict(spec, rows={'header_contains': ID_H, 'count': 200, 'number_col': 0})
            with open(sp, 'w', encoding='utf-8') as f:
                json.dump(huge, f, ensure_ascii=False)
            rc3, js3, err3 = run('--template', tpl, '--spec', sp, '--out', os.path.join(tmp, 'out', 'huge'))
            check('fit_one_page: 200 rows cannot fit: exit 1, "does not fit", row height at the 220 minimum', rc3 == 1 and 'does not fit' in err3 and js3.get('row_height') == 220 and not js3.get('ok'), (rc3, js3, err3[-300:]))
            with open(sp, 'w', encoding='utf-8') as f:
                json.dump(spec, f, ensure_ascii=False)
        else:
            print('  skipped: PDF checks (no soffice)')
        # the checks must break: an output that lost the logo, or the form number, fails with a clear message
        m = load_tool()
        orig = m.write_docx

        def lose_logo(src, dst, px):
            orig(src, dst, px)
            tmpd = dst + '.x'
            strip_images(dst, tmpd)
            os.replace(tmpd, dst)
        m.write_docx = lose_logo
        o, e = io.StringIO(), io.StringIO()
        with redirect_stdout(o), redirect_stderr(e):
            rc4 = m.main(['--template', tpl, '--spec', sp, '--out', os.path.join(tmp, 'out', 'nologo')])
        j4 = json.loads(o.getvalue().strip().splitlines()[-1])
        check('verification breaks when the output loses the logo: exit 1, "logo lost", ok false', rc4 == 1 and 'logo lost' in e.getvalue() and j4.get('ok') is False and j4.get('images') == 0, (rc4, e.getvalue()))

        def lose_form_no(src, dst, px):
            for n, root in px.items():
                if 'header' in n:
                    for t in root.iter(qn('w:t')):
                        t.text = (t.text or '').replace('08.01', '')
            orig(src, dst, px)
        m.write_docx = lose_form_no
        o, e = io.StringIO(), io.StringIO()
        with redirect_stdout(o), redirect_stderr(e):
            rc5 = m.main(['--template', tpl, '--spec', sp, '--out', os.path.join(tmp, 'out', 'nofn')])
        check('verification breaks when the form number is gone from the header: exit 1, "form number 08.01 missing"', rc5 == 1 and 'form number 08.01 missing' in e.getvalue(), (rc5, e.getvalue()))
        m.write_docx = orig
        # a template without a logo is not a reason to fail by itself (nothing lost), but a missing label is
        nologo = os.path.join(tmp, 'nologo.docx')
        strip_images(tpl, nologo)
        rc6, js6, _ = run('--template', nologo, '--spec', sp, '--out', os.path.join(tmp, 'out', 'n2'))
        check('a template with no image: images 0 in the summary (compared with the template, not a fixed number)', rc6 == 0 and js6.get('images') == 0, js6)
        bad = dict(spec, cells=[{'label': 'no such label', 'value': 'x'}])
        with open(sp, 'w', encoding='utf-8') as f:
            json.dump(bad, f, ensure_ascii=False)
        rc7, js7, err7 = run('--template', tpl, '--spec', sp, '--out', os.path.join(tmp, 'out', 'bad'))
        check('a label that is not in the form: exit 1, "label not found", no silent skip', rc7 == 1 and 'label not found' in err7 and js7.get('ok') is False, (rc7, err7))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print('\nform-fill-test: %d passed, %d failed' % (passed, len(fails)))
    sys.exit(1 if fails else 0)


if __name__ == '__main__':
    main()
